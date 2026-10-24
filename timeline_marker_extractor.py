
import os
import queue
import re
import sys
import threading
import tkinter as tk

from collections import OrderedDict
from copy import copy
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

from openpyxl import Workbook, load_workbook


# SETTINGS AND COLOURS

OUTPUT_FOLDER = "Extracted Marked"
YELLOW_RGB = "FFFF00"

BG = "#F2F5FA"
NAVY = "#15253D"
BLUE = "#2563EB"
MUTED = "#617188"
GREEN = "#15803D"


# IDENTIFY HIGHLIGHTED ROWS

def is_yellow(cell):
    fill = cell.fill

    return (
        fill.patternType == "solid"
        and fill.fgColor.type == "rgb"
        and str(fill.fgColor.rgb).upper().endswith(
            YELLOW_RGB
        )
    )


def safe_name(name):
    clean = re.sub(
        r'[<>:"/\\|?*\x00-\x1f]',
        "_",
        str(name)
    ).strip(" .")

    return clean[:110] or "Unnamed event"


def file_kind(path):
    name = path.stem.lower()

    if "audio" in name and "marked" in name:
        return "audio"

    if "detailed" in name and "marked" in name:
        return "detailed"

    raise ValueError(
        f"Not a recognised marked file: {path.name}"
    )


# FIND HEADERS AND GROUP EVENTS

def find_header(ws):

    for row in ws.iter_rows(
        min_row=1,
        max_row=min(ws.max_row, 50)
    ):
        for cell in row:
            if (
                str(cell.value or "")
                .strip()
                .lower() == "timestamp_s"
            ):
                return cell.row

    raise ValueError(
        "No timestamp_s header found."
    )


def group_events(ws):

    header = find_header(ws)
    label_column = ws.max_column

    events = OrderedDict()

    current = []
    active = False
    unknown = 0

    for row in range(
        header + 1,
        ws.max_row + 1
    ):

        if not is_yellow(ws.cell(row, 1)):
            current = []
            active = False
            continue

        label = ws.cell(
            row,
            label_column
        ).value

        if isinstance(label, str) and label.strip():

            current = [
                part.strip()
                for part in label.split("|")
                if part.strip()
            ]

        elif not active:

            unknown += 1

            current = [
                f"Unlabelled event {unknown}"
            ]

        active = True

        for name in current:
            events.setdefault(
                name, []
            ).append(row)

    return header, events


# SAFE EXCEL COPYING

def copy_row(
    source,
    target,
    source_row,
    target_row,
    width
):

    for column in range(1, width + 1):

        old = source.cell(
            source_row,
            column
        )

        new = target.cell(
            target_row,
            column,
            old.value
        )

        # copy formatting properties separately
        # never copy internal _style IDs directly

        if old.has_style:

            new.font = copy(old.font)
            new.fill = copy(old.fill)
            new.border = copy(old.border)
            new.alignment = copy(old.alignment)
            new.protection = copy(old.protection)
            new.number_format = old.number_format

        if old.hyperlink:
            new.hyperlink = copy(old.hyperlink)

        if old.comment:
            new.comment = copy(old.comment)

# CREATE EVENT WORKBOOK

def write_event(
    source,
    header,
    rows,
    destination
):

    book = Workbook()

    sheet = book.active
    sheet.title = "Marked Data"

    width = source.max_column

    source_rows = (
        list(range(1, header + 1))
        + rows
    )

    for output_row, source_row in enumerate(
        source_rows,
        start=1
    ):

        copy_row(
            source,
            sheet,
            source_row,
            output_row,
            width
        )

    # preserve column widths

    for key, dimension in (
        source.column_dimensions.items()
    ):

        sheet.column_dimensions[key].width = (
            dimension.width
        )

        sheet.column_dimensions[key].hidden = (
            dimension.hidden
        )

    sheet.freeze_panes = f"A{header + 1}"

    destination.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    temp = destination.with_name(
        destination.stem + ".working.xlsx"
    )

    try:

        book.save(temp)

        # verify the workbook can be read
        # before replacing the output file

        check = load_workbook(
            temp,
            read_only=True,
            data_only=True
        )

        check.close()

        os.replace(
            temp,
            destination
        )

    finally:

        book.close()

        if temp.exists():
            temp.unlink()


# MAIN EXTRACTION ENGINE

def extract(paths, notify):

    workbooks = []
    prepared = []

    try:

        # STAGE 1: SCAN INPUT FILES

        for path in paths:

            kind = file_kind(path)

            notify(
                "scan",
                f"Reading {path.name}"
            )

            wb = load_workbook(
                path,
                data_only=False,
                read_only=False
            )

            workbooks.append(wb)

            ws = wb.worksheets[0]

            header, events = group_events(ws)

            if not events:

                raise ValueError(
                    "No yellow-highlighted events "
                    f"found in {path.name}."
                )

            prepared.append((
                kind,
                ws,
                header,
                events
            ))

        # calculate total number of output files

        labels = {
            kind: set(events)
            for kind, _, _, events in prepared
        }

        total = sum(
            len(events)
            for _, _, _, events in prepared
        )

        notify("plan", total)

        destination = (
            paths[0].parent /
            OUTPUT_FOLDER
        )

        completed = 0

        # STAGE 2: EXPORT EVENT FILES

        for kind, ws, header, events in prepared:

            for event, rows in events.items():

                if kind == "audio":

                    stem = "audio_features_marked_"

                else:

                    stem = "input_detailed_marked_"

                filename = (
                    f"{stem}"
                    f"({safe_name(event)}).xlsx"
                )

                write_event(
                    ws,
                    header,
                    rows,
                    destination / filename
                )

                completed += 1

                # send actual progress to the UI

                notify(
                    "progress",
                    (
                        completed,
                        total,
                        filename,
                        len(rows)
                    )
                )

        warnings = []

        if labels["audio"] != labels["detailed"]:

            warnings.append(
                "Audio and Detailed contain "
                "different event labels."
            )

        notify(
            "done",
            (
                destination,
                completed,
                warnings
            )
        )

    except Exception as error:

        notify("error", str(error))

    finally:

        for wb in workbooks:
            wb.close()


# APPLICATION WINDOW


class ExtractorApp(tk.Tk):

    def __init__(self, arguments=()):

        super().__init__()

        self.title(
            "Highlighted Timeline Extractor"
        )

        self.geometry("820x700")
        self.minsize(720, 650)

        self.configure(bg=BG)

        self.detailed = tk.StringVar()
        self.audio = tk.StringVar()

        self.events = queue.Queue()

        self.running = False
        self.output = None

        self.completed = 0
        self.total = 0

        self.input_controls = []

        self.make_styles()
        self.make_window()

        self.after(100, self.poll)

        self.protocol(
            "WM_DELETE_WINDOW",
            self.close_window
        )

        # support windows send to

        if arguments:

            for raw in arguments:

                p = (
                    Path(raw)
                    .expanduser()
                    .resolve()
                )

                kind = file_kind(p)

                if kind == "audio":
                    self.audio.set(str(p))
                else:
                    self.detailed.set(str(p))

            self.after(250, self.start)


# UI STYLES

    def make_styles(self):

        style = ttk.Style(self)

        style.theme_use("clam")

        style.configure(
            "Blue.Horizontal.TProgressbar",
            troughcolor="#DFE6F0",
            background=BLUE,
            bordercolor="#DFE6F0",
            lightcolor=BLUE,
            darkcolor=BLUE,
            thickness=18
        )

        style.configure(
            "App.TEntry",
            padding=8,
            font=("Segoe UI", 10)
        )


# BUILD INTERFACE

    def make_window(self):

        # TOP HEADER

        header = tk.Frame(
            self,
            bg=NAVY,
            padx=30,
            pady=16
        )

        header.pack(fill="x")

        tk.Label(
            header,
            text="HIGHLIGHTED TIMELINE EXTRACTOR",
            bg=NAVY,
            fg="white",
            font=("Segoe UI", 19, "bold"),
            anchor="w"
        ).pack(fill="x")

        tk.Label(
            header,
            text=(
                "Create a separate Excel file "
                "for every highlighted event"
            ),
            bg=NAVY,
            fg="#CAD5E3",
            font=("Segoe UI", 10),
            anchor="w"
        ).pack(
            fill="x",
            pady=(5, 0)
        )

        main = tk.Frame(
            self,
            bg=BG,
            padx=26,
            pady=12
        )

        main.pack(
            fill="both",
            expand=True
        )

        # FILE SELECTION PANEL

        files = tk.Frame(
            main,
            bg="white",
            padx=18,
            pady=12,
            highlightbackground="#DFE6F0",
            highlightthickness=1
        )

        files.pack(fill="x")

        tk.Label(
            files,
            text="1   SELECT YOUR MARKED FILES",
            bg="white",
            fg=NAVY,
            font=("Segoe UI", 11, "bold"),
            anchor="w"
        ).pack(
            fill="x",
            pady=(0, 6)
        )

        self.file_row(
            files,
            "Detailed data",
            self.detailed
        )

        self.file_row(
            files,
            "Audio features",
            self.audio
        )

        # ACTION BUTTONS

        actions = tk.Frame(
            main,
            bg=BG
        )

        actions.pack(
            fill="x",
            pady=(10, 8)
        )

        self.start_button = tk.Button(
            actions,
            text="Extract highlighted events",
            command=self.start,
            bg=BLUE,
            fg="white",
            activebackground="#1D4ED8",
            activeforeground="white",
            font=("Segoe UI", 11, "bold"),
            padx=17,
            pady=10,
            relief="flat",
            cursor="hand2"
        )

        self.start_button.pack(side="left")

        self.open_button = tk.Button(
            actions,
            text="Open results folder",
            command=self.open_results,
            bg="white",
            fg=NAVY,
            font=("Segoe UI", 10),
            padx=15,
            pady=10,
            relief="flat",
            state="disabled"
        )

        self.open_button.pack(side="right")

        # PROGRESS PANEL

        progress_area = tk.Frame(
            main,
            bg="white",
            padx=18,
            pady=10,
            highlightbackground="#DFE6F0",
            highlightthickness=1
        )

        progress_area.pack(fill="x")

        title_line = tk.Frame(
            progress_area,
            bg="white"
        )

        title_line.pack(fill="x")

        tk.Label(
            title_line,
            text="2   EXTRACTION PROGRESS",
            bg="white",
            fg=NAVY,
            font=("Segoe UI", 11, "bold")
        ).pack(side="left")

        self.percent_label = tk.Label(
            title_line,
            text="0%",
            bg="white",
            fg=BLUE,
            font=("Segoe UI", 13, "bold")
        )

        self.percent_label.pack(side="right")

        self.bar = ttk.Progressbar(
            progress_area,
            mode="determinate",
            maximum=100,
            value=0,
            style="Blue.Horizontal.TProgressbar"
        )

        self.bar.pack(
            fill="x",
            pady=(7, 5)
        )

        self.status_label = tk.Label(
            progress_area,
            text="Ready to extract",
            bg="white",
            fg=MUTED,
            anchor="w",
            font=("Segoe UI", 10)
        )

        self.status_label.pack(fill="x")

        self.count_label = tk.Label(
            progress_area,
            text="0 of 0 files completed",
            bg="white",
            fg=MUTED,
            anchor="w",
            font=("Segoe UI", 9)
        )

        self.count_label.pack(
            fill="x",
            pady=(4, 0)
        )

        # ACTIVITY LOG

        log_header = tk.Frame(
            main,
            bg=BG
        )

        log_header.pack(
            fill="x",
            pady=(10, 4)
        )

        tk.Label(
            log_header,
            text="ACTIVITY LOG",
            bg=BG,
            fg=NAVY,
            font=("Segoe UI", 10, "bold")
        ).pack(side="left")

        tk.Label(
            log_header,
            text=f"Output: {OUTPUT_FOLDER}",
            bg=BG,
            fg=MUTED,
            font=("Segoe UI", 9)
        ).pack(side="right")

        log_wrap = tk.Frame(
            main,
            bg="white"
        )

        log_wrap.pack(
            fill="both",
            expand=True
        )

        self.log = tk.Text(
            log_wrap,
            height=6,
            font=("Consolas", 9),
            bg="white",
            fg="#334155",
            bd=0,
            padx=12,
            pady=10,
            state="disabled",
            wrap="word"
        )

        self.log.pack(
            side="left",
            fill="both",
            expand=True
        )

        scroll = ttk.Scrollbar(
            log_wrap,
            command=self.log.yview
        )

        scroll.pack(
            side="right",
            fill="y"
        )

        self.log.configure(
            yscrollcommand=scroll.set
        )

# FILE SELECTION CONTROLS

    def file_row(
        self,
        parent,
        caption,
        variable
    ):

        tk.Label(
            parent,
            text=caption,
            bg="white",
            fg="#3F516A",
            font=("Segoe UI", 10, "bold")
        ).pack(
            fill="x",
            anchor="w"
        )

        row = tk.Frame(
            parent,
            bg="white"
        )

        row.pack(
            fill="x",
            pady=(3, 6)
        )

        entry = ttk.Entry(
            row,
            textvariable=variable,
            style="App.TEntry"
        )

        entry.pack(
            side="left",
            fill="x",
            expand=True
        )

        button = tk.Button(
            row,
            text="Browse",
            command=lambda: self.browse(variable),
            bg="#E7EDF7",
            fg=NAVY,
            padx=16,
            pady=7,
            font=("Segoe UI", 9, "bold"),
            relief="flat"
        )

        button.pack(
            side="left",
            padx=(10, 0)
        )

        self.input_controls.append(
            (entry, button)
        )


    def set_input_enabled(self, enabled):

        for entry, button in self.input_controls:

            state = (
                "normal"
                if enabled
                else "disabled"
            )

            entry.configure(state=state)
            button.configure(state=state)


    def browse(self, variable):

        if self.running:
            return

        selected = filedialog.askopenfilename(
            title="Select marked Excel file",
            filetypes=[
                ("Excel files", "*.xlsx")
            ]
        )

        if selected:
            variable.set(selected)


# ACTIVITY LOG

    def add_log(self, line):

        self.log.configure(
            state="normal"
        )

        self.log.insert(
            "end",
            line + "\n"
        )

        self.log.see("end")

        self.log.configure(
            state="disabled"
        )


# START EXTRACTION

    def start(self):

        if self.running:
            return

        raw_paths = [
            self.detailed.get().strip(),
            self.audio.get().strip()
        ]

        if not all(raw_paths):

            messagebox.showwarning(
                "Files needed",
                "Select both marked Excel files first."
            )

            return

        paths = [
            Path(p).expanduser().resolve()
            for p in raw_paths
        ]

        try:

            if any(
                not p.is_file()
                or p.suffix.lower() != ".xlsx"
                for p in paths
            ):

                raise ValueError(
                    "Both files must be existing "
                    ".xlsx workbooks."
                )

            if len({
                p.parent for p in paths
            }) != 1:

                raise ValueError(
                    "Both files must be in the same folder."
                )

            kinds = [
                file_kind(p)
                for p in paths
            ]

            if kinds != ["detailed", "audio"]:

                raise ValueError(
                    "Choose one Detailed and one "
                    "Audio Features marked file."
                )

            output = (
                paths[0].parent /
                OUTPUT_FOLDER
            )

            if (
                output.is_dir()
                and any(output.glob("*.xlsx"))
            ):

                confirmed = messagebox.askyesno(
                    "Existing extracted files",
                    "Matching output files will be "
                    "replaced. Continue?"
                )

                if not confirmed:
                    return

        except ValueError as error:

            messagebox.showerror(
                "Check your files",
                str(error)
            )

            return

        # RESET USER INTERFACE

        self.running = True
        self.output = None

        self.total = 0
        self.completed = 0

        self.start_button.configure(
            state="disabled"
        )

        self.open_button.configure(
            state="disabled"
        )

        self.set_input_enabled(False)

        self.percent_label.configure(
            text="0%",
            fg=BLUE
        )

        self.count_label.configure(
            text="Scanning source files..."
        )

        # Animated progress while detecting events

        self.bar.configure(
            mode="indeterminate",
            value=0
        )

        self.bar.start(12)

        self.status_label.configure(
            text="Scanning highlighted events..."
        )

        self.add_log(
            "-- Started new extraction --"
        )

        # heavy processing runs in another thread
        # keeping the window responsive

        threading.Thread(
            target=extract,
            args=(paths, self.send),
            daemon=True
        ).start()



# THREAD SAFE PROGRESS UPDATES


    def send(self, kind, data):

        self.events.put(
            (kind, data)
        )


    def poll(self):

        try:

            while True:

                kind, data = (
                    self.events.get_nowait()
                )

                # READING FILES

                if kind == "scan":

                    self.status_label.configure(
                        text=data
                    )

                    self.add_log(data)

                # FOUND TOTAL FILE COUNT

                elif kind == "plan":

                    self.total = data

                    self.bar.stop()

                    # switch to real progress tracking

                    self.bar.configure(
                        mode="determinate",
                        maximum=data,
                        value=0
                    )

                    self.count_label.configure(
                        text=f"0 of {data} files completed"
                    )

                    self.status_label.configure(
                        text="Exporting event workbooks..."
                    )

                    self.add_log(
                        f"Found {data} event files to create."
                    )

                # FILE SAVED SUCCESSFULLY

                elif kind == "progress":

                    (
                        self.completed,
                        total,
                        name,
                        rows
                    ) = data

                    self.bar.configure(
                        value=self.completed
                    )

                    percent = round(
                        self.completed * 100 / total
                    )

                    self.percent_label.configure(
                        text=f"{percent}%"
                    )

                    self.count_label.configure(
                        text=(
                            f"{self.completed} of "
                            f"{total} files completed"
                        )
                    )

                    self.status_label.configure(
                        text=f"Saved: {name}"
                    )

                    self.add_log(
                        f"[{self.completed}/{total}] "
                        f"{name} ({rows} rows)"
                    )

                # ALL FILES COMPLETED

                elif kind == "done":

                    (
                        destination,
                        count,
                        warnings
                    ) = data

                    self.running = False
                    self.output = destination

                    self.bar.stop()

                    self.bar.configure(
                        mode="determinate",
                        maximum=max(count, 1),
                        value=count
                    )

                    self.percent_label.configure(
                        text="100%",
                        fg=GREEN
                    )

                    self.status_label.configure(
                        text="Extraction completed successfully"
                    )

                    self.count_label.configure(
                        text=(
                            f"{count} of {count} "
                            "files completed"
                        )
                    )

                    self.start_button.configure(
                        state="normal"
                    )

                    self.open_button.configure(
                        state="normal"
                    )

                    self.set_input_enabled(True)

                    self.add_log(
                        f"SUCCESS: {count} Excel files "
                        f"saved to {destination}"
                    )

                    for warning in warnings:
                        self.add_log(
                            "NOTE: " + warning
                        )

                    messagebox.showinfo(
                        "Extraction complete",
                        f"Created {count} Excel files.\n\n"
                        f"Saved in:\n{destination}"
                    )

                # PROCESSING ERROR

                elif kind == "error":

                    self.running = False

                    self.bar.stop()

                    self.bar.configure(
                        mode="determinate"
                    )

                    self.start_button.configure(
                        state="normal"
                    )

                    self.set_input_enabled(True)

                    self.percent_label.configure(
                        fg=BLUE
                    )

                    self.status_label.configure(
                        text=(
                            "Extraction stopped "
                            "due to an error"
                        )
                    )

                    self.add_log(
                        "ERROR: " + data
                    )

                    messagebox.showerror(
                        "Extraction failed",
                        data
                    )

        except queue.Empty:
            pass

        self.after(
            100,
            self.poll
        )



# OPEN RESULTS AND CLOSE WINDOW

    def close_window(self):

        if self.running:

            messagebox.showinfo(
                "Extraction running",
                "Finish extracting before closing "
                "this window."
            )

            return

        self.destroy()


    def open_results(self):

        if self.output and self.output.is_dir():

            os.startfile(
                str(self.output)
            )

        else:

            messagebox.showinfo(
                "No results",
                "Run the extractor first."
            )


# START APPLICATION

if __name__ == "__main__":

    args = sys.argv[1:]

    if args and len(args) != 2:

        temp = tk.Tk()
        temp.withdraw()

        messagebox.showerror(
            "Select two files",
            "Select exactly two marked Excel files."
        )

        temp.destroy()

        sys.exit(2)

    try:

        ExtractorApp(args).mainloop()

    except Exception as error:

        temp = tk.Tk()
        temp.withdraw()

        messagebox.showerror(
            "Timeline Extractor",
            str(error)
        )

        temp.destroy()
