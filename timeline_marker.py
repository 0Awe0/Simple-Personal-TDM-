
import csv
import os
import queue
import re
import threading
import tkinter as tk
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import time, timedelta
from decimal import Decimal
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

YELLOW = PatternFill(fill_type="solid", fgColor="FFFFFF00")


# TIMESTAMP PARSING

def seconds(value):
    if isinstance(value, timedelta):
        return Decimal(str(value.total_seconds()))
    if isinstance(value, time):
        return (Decimal(value.hour * 3600 + value.minute * 60 + value.second)
                + Decimal(value.microsecond) / 1000000)

    text = str(value).strip().replace(",", ".")
    parts = re.split(r"[:\-]", text)

    if not 1 <= len(parts) <= 3:
        raise ValueError(f"Invalid timestamp: {value}")

    numbers = [Decimal(p.strip()) for p in parts]

    if len(numbers) > 1 and numbers[-1] >= 60:
        raise ValueError(f"Invalid seconds: {value}")
    if len(numbers) == 3 and numbers[-2] >= 60:
        raise ValueError(f"Invalid minutes: {value}")

    return sum(n * 60 ** i for i, n in enumerate(reversed(numbers)))


def read_rows(path):
    if path.suffix.lower() == ".csv":
        with path.open(
            encoding="utf-8-sig",
            newline="",
            errors="replace"
        ) as f:
            yield from csv.reader(f)
    else:
        book = load_workbook(
            path, read_only=True, data_only=True
        )
        try:
            yield from book.worksheets[0].iter_rows(
                values_only=True
            )
        finally:
            book.close()


def read_timeline(path):
    periods = []

    for row_number, row in enumerate(read_rows(path), 1):
        if not row or row[0] is None or not str(row[0]).strip():
            continue

        name = str(row[0]).strip()

        if name.lower() in ("name", "label", "timeline", "event"):
            continue

        if len(row) < 2 or row[1] is None:
            raise ValueError(
                f"Timeline row {row_number}: missing time range"
            )

        if len(row) >= 3 and row[2] not in (None, ""):
            start, end = row[1], row[2]
        else:
            parts = re.split(
                r"\s+to\s+",
                str(row[1]).strip(),
                maxsplit=1,
                flags=re.I
            )
            if len(parts) != 2:
                raise ValueError(
                    f"Timeline row {row_number}: use START to END"
                )
            start, end = parts

        first, last = seconds(start), seconds(end)

        if last < first:
            raise ValueError(
                f"Timeline row {row_number}: end before start"
            )

        periods.append((name, first, last))

    if not periods:
        raise ValueError("No timeline intervals found")

    return periods


# EXCEL MARKING

def inspect_data(path):
    width, header_row, timestamp_col = 0, None, None

    for index, row in enumerate(read_rows(path)):
        width = max(width, len(row))

        if index < 50 and header_row is None:
            for col, value in enumerate(row):
                if str(value or "").strip().lower() == "timestamp_s":
                    header_row, timestamp_col = index, col
                    break

    if header_row is None:
        raise ValueError(
            f"{path.name}: timestamp_s column not found"
        )

    if width >= 16384:
        raise ValueError(f"{path.name}: too many columns")

    return width, header_row, timestamp_col


def highlight_file(source, periods, destination):
    width, header_row, timestamp_col = inspect_data(source)

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Marked Data")

    sheet.freeze_panes = f"A{header_row + 2}"
    sheet.column_dimensions[
        get_column_letter(width + 1)
    ].width = 23

    seen = set()
    total = 0

    for row_number, record in enumerate(read_rows(source)):
        row = list(record)
        matching = []

        if row_number > header_row and timestamp_col < len(row):
            try:
                stamp = seconds(row[timestamp_col])
                matching = [
                    i for i, (_, start, end) in enumerate(periods)
                    if start <= stamp <= end
                ]
            except (ValueError, TypeError, ArithmeticError):
                pass

        if matching:
            total += 1

            labels = [
                periods[i][0] for i in matching
                if i not in seen
            ]

            seen.update(matching)

            painted = []

            for col in range(width):
                cell = WriteOnlyCell(
                    sheet,
                    value=row[col] if col < len(row) else None
                )
                cell.fill = YELLOW
                painted.append(cell)

            if labels:
                tag = WriteOnlyCell(
                    sheet,
                    value=" | ".join(labels)
                )
                tag.fill = YELLOW
                tag.font = Font(bold=True)
                painted.append(tag)

            sheet.append(painted)
        else:
            sheet.append(row)

    destination.parent.mkdir(
        parents=True, exist_ok=True
    )

    temporary = destination.with_name(
        destination.stem + ".working.xlsx"
    )

    try:
        workbook.save(temporary)
        os.replace(temporary, destination)
    finally:
        if temporary.exists():
            temporary.unlink()

    return total, len(periods) - len(seen)


# BULK DETECTION

def inspect_sets(root):
    all_sets = []

    folders = [root] + sorted(
        p for p in root.iterdir() if p.is_dir()
    )

    for folder in folders:
        names = [
            p for p in folder.iterdir()
            if p.is_file()
            and p.suffix.lower() in (".csv", ".xlsx")
            and not p.name.startswith("~$")
            and "_marked" not in p.stem.lower()
        ]

        timelines = [
            p for p in names
            if p.suffix.lower() == ".xlsx"
            and "timeline" in p.stem.lower()
        ]

        details = [
            p for p in names
            if "detailed" in p.stem.lower()
        ]

        audios = [
            p for p in names
            if re.search(
                r"audio[_ -]?features",
                p.stem, re.I
            )
        ]

        if not (timelines or details or audios):
            continue

        valid = (
            len(timelines),
            len(details),
            len(audios)
        ) == (1, 1, 1)

        status = "Ready" if valid else (
            "Needs 1 timeline, 1 detailed, 1 audio "
            f"(found {len(timelines)}/"
            f"{len(details)}/{len(audios)})"
        )

        all_sets.append((
            folder, timelines, details,
            audios, valid, status
        ))

    return all_sets


def process_set(entry, out):
    folder, timelines, details, audios, _, _ = entry

    periods = read_timeline(timelines[0])
    destination = out / folder.name

    d, d_missed = highlight_file(
        details[0],
        periods,
        destination / "detailed_marked.xlsx"
    )

    a, a_missed = highlight_file(
        audios[0],
        periods,
        destination / "audio_features_marked.xlsx"
    )

    return folder.name, d, a, d_missed, a_missed


# UI PART

class TimelineApp(tk.Tk):
    def __init__(self):
        super().__init__()

        self.title("Timeline Marker | Bulk Highlighter")
        self.geometry("880x660")
        self.minsize(720, 550)
        self.configure(bg="#F3F6FA")

        self.input_dir = tk.StringVar()
        self.output_dir = tk.StringVar()
        self.events = queue.Queue()
        self.sets = []
        self.busy = False

        style = ttk.Style(self)
        style.theme_use("clam")

        style.configure(
            "Treeview",
            font=("Segoe UI", 10),
            rowheight=29
        )

        style.configure(
            "Treeview.Heading",
            font=("Segoe UI", 10, "bold")
        )

        style.configure(
            "App.Horizontal.TProgressbar",
            troughcolor="#DBEAFE",
            background="#2563EB"
        )

        self.build()
        self.after(100, self.poll)

    def build(self):
        body = tk.Frame(
            self, bg="#F3F6FA",
            padx=24, pady=18
        )
        body.pack(fill="both", expand=True)

        tk.Label(
            body,
            text="Timeline Marker",
            font=("Segoe UI", 23, "bold"),
            bg="#F3F6FA",
            fg="#14253D"
        ).pack(anchor="w")

        tk.Label(
            body,
            text="Highlight matching timestamps "
                 "across multiple recording sets in one click.",
            font=("Segoe UI", 10),
            bg="#F3F6FA",
            fg="#52657D"
        ).pack(anchor="w", pady=(0, 18))

        panel = tk.Frame(
            body, bg="white",
            padx=18, pady=16
        )
        panel.pack(fill="x")

        self.path_row(
            panel,
            "Folder containing your sets",
            self.input_dir,
            self.pick_input
        )

        self.path_row(
            panel,
            "Where to save results",
            self.output_dir,
            self.pick_output
        )

        actions = tk.Frame(
            body, bg="#F3F6FA"
        )
        actions.pack(fill="x", pady=(17, 11))

        self.scan_btn = tk.Button(
            actions,
            text="1  Scan folders",
            command=self.scan,
            bg="#E4EAF3",
            fg="#172B4D",
            bd=0,
            padx=20,
            pady=11,
            font=("Segoe UI", 10, "bold"),
            cursor="hand2"
        )
        self.scan_btn.pack(side="left")

        self.run_btn = tk.Button(
            actions,
            text="2  Process all ready sets",
            command=self.run_batch,
            bg="#2563EB",
            fg="white",
            bd=0,
            padx=20,
            pady=11,
            font=("Segoe UI", 10, "bold"),
            cursor="hand2"
        )
        self.run_btn.pack(side="left", padx=10)

        tk.Button(
            actions,
            text="Open results",
            command=self.open_results,
            bg="white",
            fg="#172B4D",
            bd=0,
            padx=18,
            pady=11,
            font=("Segoe UI", 10),
            cursor="hand2"
        ).pack(side="right")

        heading = tk.Frame(
            body, bg="#F3F6FA"
        )
        heading.pack(fill="x", pady=(4, 6))

        tk.Label(
            heading,
            text="Detected sets",
            font=("Segoe UI", 12, "bold"),
            bg="#F3F6FA",
            fg="#14253D"
        ).pack(side="left")

        self.count = tk.Label(
            heading,
            text="0 ready",
            font=("Segoe UI", 10),
            bg="#F3F6FA",
            fg="#52657D"
        )
        self.count.pack(side="right")

        frame = ttk.Frame(body)
        frame.pack(fill="both", expand=True)

        self.table = ttk.Treeview(
            frame,
            columns=("name", "status"),
            show="headings",
            height=7
        )

        self.table.heading(
            "name", text="Set / Folder"
        )
        self.table.heading(
            "status", text="Status"
        )
        self.table.column(
            "name", width=260
        )
        self.table.column(
            "status", width=490
        )

        self.table.pack(
            side="left", fill="both", expand=True
        )

        scrollbar = ttk.Scrollbar(
            frame,
            orient="vertical",
            command=self.table.yview
        )
        scrollbar.pack(side="right", fill="y")

        self.table.configure(
            yscrollcommand=scrollbar.set
        )

        self.table.tag_configure(
            "ready", foreground="#127B49"
        )
        self.table.tag_configure(
            "warning", foreground="#AB5A10"
        )

        self.status = tk.Label(
            body,
            text="Ready to scan",
            anchor="w",
            bg="#F3F6FA",
            fg="#52657D",
            font=("Segoe UI", 10)
        )
        self.status.pack(
            fill="x", pady=(15, 5)
        )

        self.progress = ttk.Progressbar(
            body,
            mode="determinate",
            maximum=100,
            style="App.Horizontal.TProgressbar"
        )
        self.progress.pack(fill="x")

        self.logs = tk.Text(
            body,
            height=5,
            font=("Consolas", 9),
            state="disabled",
            bg="white",
            fg="#334155",
            relief="flat",
            padx=12,
            pady=10
        )
        self.logs.pack(
            fill="x", pady=(12, 0)
        )

    def path_row(self, parent, caption, variable, action):
        tk.Label(
            parent,
            text=caption,
            bg="white",
            fg="#334155",
            font=("Segoe UI", 10, "bold")
        ).pack(anchor="w", pady=(0, 4))

        row = tk.Frame(
            parent, bg="white"
        )
        row.pack(fill="x", pady=(0, 11))

        ttk.Entry(
            row,
            textvariable=variable
        ).pack(
            side="left",
            fill="x",
            expand=True,
            ipady=5
        )

        tk.Button(
            row,
            text="Browse...",
            command=action,
            bg="#E4EAF3",
            fg="#14253D",
            bd=0,
            padx=14,
            pady=7
        ).pack(side="left", padx=(10, 0))

    def pick_input(self):
        path = filedialog.askdirectory(
            title="Choose the folder with your recording sets"
        )

        if path:
            self.input_dir.set(path)

            if not self.output_dir.get():
                self.output_dir.set(
                    str(
                        Path(path).parent /
                        "Timeline_Results"
                    )
                )

            self.scan()

    def pick_output(self):
        path = filedialog.askdirectory(
            title="Choose where to save results"
        )

        if path:
            self.output_dir.set(path)

    def log(self, text):
        self.logs.configure(state="normal")
        self.logs.insert(
            "end", text + "\n"
        )
        self.logs.see("end")
        self.logs.configure(state="disabled")

    def scan(self):
        if self.busy:
            return

        path = Path(self.input_dir.get())

        if not path.is_dir():
            messagebox.showwarning(
                "Choose a folder",
                "Select the folder containing your sets first."
            )
            return

        self.sets = inspect_sets(path)

        for item in self.table.get_children():
            self.table.delete(item)

        for folder, *_rest, valid, status in self.sets:
            self.table.insert(
                "",
                "end",
                values=(folder.name, status),
                tags=("ready" if valid else "warning",)
            )

        ready = sum(
            entry[4] for entry in self.sets
        )

        self.count.configure(
            text=f"{ready} ready / {len(self.sets)} detected"
        )

        self.status.configure(
            text=f"Found {ready} complete set(s)."
        )

        self.log(
            f"Scan completed: {ready} ready, "
            f"{len(self.sets) - ready} need attention."
        )

    def run_batch(self):
        if self.busy:
            return

        ready = [
            entry for entry in self.sets
            if entry[4]
        ]

        if not ready:
            messagebox.showwarning(
                "No sets ready",
                "Scan folders and check your files first."
            )
            return

        folder = self.output_dir.get().strip()

        if not folder:
            messagebox.showwarning(
                "Choose output folder",
                "Select where to save results first."
            )
            return

        destination = Path(folder)

        prior = sum(
            (
                destination / job[0].name /
                "detailed_marked.xlsx"
            ).exists()
            or (
                destination / job[0].name /
                "audio_features_marked.xlsx"
            ).exists()
            for job in ready
        )

        if prior:
            confirmed = messagebox.askyesno(
                "Replace existing outputs?",
                f"{prior} set(s) already have output files. "
                "Do you want to overwrite them?"
            )
            if not confirmed:
                return

        self.busy = True

        self.scan_btn.configure(
            state="disabled"
        )
        self.run_btn.configure(
            state="disabled"
        )

        self.progress.configure(
            maximum=len(ready),
            value=0
        )

        self.status.configure(
            text=f"Processing 0 of {len(ready)} sets..."
        )

        def work():
            with ThreadPoolExecutor(
                max_workers=min(3, len(ready))
            ) as executor:

                jobs = {
                    executor.submit(
                        process_set,
                        entry,
                        destination
                    ): entry[0].name
                    for entry in ready
                }

                for job in as_completed(jobs):
                    try:
                        name, detailed, audio, missed_d, missed_a = (
                            job.result()
                        )

                        self.events.put((
                            "log",
                            f"SUCCESS  {name} | "
                            f"Detailed {detailed}, "
                            f"Audio {audio} highlighted; "
                            f"unmatched intervals "
                            f"{missed_d}/{missed_a}"
                        ))

                    except Exception as exc:
                        self.events.put((
                            "log",
                            f"FAILED   {jobs[job]} | {exc}"
                        ))

                    self.events.put((
                        "step", len(ready)
                    ))

            self.events.put((
                "finished", len(ready)
            ))

        threading.Thread(
            target=work,
            daemon=True
        ).start()

    def poll(self):
        try:
            while True:
                kind, data = self.events.get_nowait()

                if kind == "log":
                    self.log(data)

                elif kind == "step":
                    self.progress.step(1)

                    self.status.configure(
                        text=(
                            f"Processed "
                            f"{int(self.progress['value'])} "
                            f"of {data} sets"
                        )
                    )

                elif kind == "finished":
                    self.busy = False

                    self.scan_btn.configure(
                        state="normal"
                    )
                    self.run_btn.configure(
                        state="normal"
                    )

                    self.status.configure(
                        text=(
                            f"Finished processing {data} sets. "
                            "Open results to see the files."
                        )
                    )

                    messagebox.showinfo(
                        "Processing complete",
                        "Your batch has finished. "
                        "See the log for any errors."
                    )

        except queue.Empty:
            pass

        self.after(100, self.poll)

    def open_results(self):
        path = Path(
            self.output_dir.get().strip()
        )

        if path.is_dir():
            os.startfile(str(path))
        else:
            messagebox.showinfo(
                "No output folder",
                "Choose and process an output folder first."
            )


if __name__ == "__main__":
    TimelineApp().mainloop()
